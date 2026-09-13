# -*- coding: utf-8 -*-
"""
Builds template/cv_template.docx by injecting docxtpl placeholders into the
ORIGINAL agency form (template/original.docx) - preserving its exact design
(logo, colours, layout, photo boxes).

Only the modern (mc:Choice) table is edited - that is what Word / LibreOffice
render. Re-run after replacing template/original.docx:
    python src/build_template.py
"""
import os
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
ORIGINAL = os.path.join(ROOT, "template", "original.docx")
OUT = os.path.join(ROOT, "template", "cv_template.docx")

# (row, cell) -> (kind, placeholder)
#   text : empty value cell  -> insert placeholder
#   rep  : cell has default text (HOUSE MAID / SAR / KG ...) -> replace it
#   img  : photo cell -> insert an image placeholder
MAP = {
    (1, 1):  ("img",  "photo_logo"),
    (1, 2):  ("img",  "photo_face"),
    (3, 2):  ("text", "full_name"),
    (4, 3):  ("rep",  "position"),
    (5, 3):  ("rep",  "salary"),
    (6, 3):  ("rep",  "preferred_country"),
    (8, 3):  ("text", "passport_number"),
    (9, 3):  ("text", "date_of_issue"),
    (10, 3): ("text", "date_of_expiry"),
    (11, 3): ("text", "place_of_issue"),
    (13, 3): ("text", "english_level"),
    (14, 3): ("text", "arabic_level"),
    (15, 3): ("text", "education_level"),
    (19, 2): ("text", "prev_period_from"),
    (19, 3): ("text", "prev_period_to"),
    (19, 4): ("text", "prev_country"),
    (19, 5): ("text", "prev_position"),
    (23, 3): ("text", "nationality"),
    (24, 3): ("text", "contact_no"),
    (25, 3): ("text", "religion"),
    (26, 3): ("text", "date_of_birth"),
    (27, 3): ("text", "place_of_birth"),
    (28, 3): ("text", "living_town"),
    (29, 3): ("text", "marital_status"),
    (30, 2): ("text", "skill_babysitting"),
    (30, 5): ("text", "num_children"),
    (31, 2): ("text", "skill_children_care"),
    (31, 5): ("rep",  "weight"),
    (32, 2): ("text", "skill_tutoring"),
    (32, 5): ("rep",  "height"),
    (33, 2): ("text", "skill_disabled_care"),
    (33, 5): ("text", "complexion"),
    (34, 2): ("text", "skill_cleaning"),
    (34, 5): ("rep",  "age"),
    (35, 2): ("text", "skill_washing"),
    (36, 2): ("text", "skill_ironing"),
    (36, 4): ("text", "profile_summary"),
    (37, 2): ("text", "skill_cooking"),
    (38, 2): ("text", "skill_arabic_cooking"),
}

# label we expect next to each value cell, used to verify the mapping is aligned
VERIFY = {
    (3, 1): "FULL NAME", (4, 2): "POSITION", (8, 2): "NUMBER",
    (23, 2): "NATIONALITY", (30, 1): "BABBY SITTING", (34, 4): "AGE",
}


def _choice_table(doc):
    for t in doc.element.body.xpath(".//w:tbl"):
        for a in t.iterancestors():
            if a.tag.split("}")[-1] == "Choice":
                return t
    # fallback: first table
    tbls = doc.element.body.xpath(".//w:tbl")
    return tbls[0] if tbls else None


def _cell_text(tc):
    return "".join(x.text or "" for x in tc.iter(qn("w:t"))).strip()


def _make_run(text, bold=False, size=18, color="000000"):
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(a), "Arial")
    rPr.append(rFonts)
    if bold:
        rPr.append(OxmlElement("w:b"))
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(size)); rPr.append(sz)
    szc = OxmlElement("w:szCs"); szc.set(qn("w:val"), str(size)); rPr.append(szc)
    col = OxmlElement("w:color"); col.set(qn("w:val"), color); rPr.append(col)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    return r


def _first_p(tc):
    p = tc.find(qn("w:p"))
    if p is None:
        p = OxmlElement("w:p")
        tc.append(p)
    return p


# elements that must appear AFTER w:vMerge inside w:tcPr (schema order)
_AFTER_VMERGE = {"tcBorders", "shd", "noWrap", "tcMar", "textDirection",
                 "tcFit", "vAlign", "hideMark", "headers"}


def _set_vmerge(tc, restart):
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        tcPr = OxmlElement("w:tcPr")
        tc.insert(0, tcPr)
    for existing in tcPr.findall(qn("w:vMerge")):
        tcPr.remove(existing)
    vm = OxmlElement("w:vMerge")
    if restart:
        vm.set(qn("w:val"), "restart")
    idx = len(tcPr)
    for i, ch in enumerate(tcPr):
        if ch.tag.split("}")[-1] in _AFTER_VMERGE:
            idx = i
            break
    tcPr.insert(idx, vm)


def _center_cell(tc):
    """Center the cell's first paragraph horizontally (used for photo cells)."""
    Paragraph(_first_p(tc), None).alignment = WD_ALIGN_PARAGRAPH.CENTER


def _format_photo_cell(tc):
    """Zero cell margins, center alignment, and clean paragraph spacing for photo containers."""
    tcPr = tc.get_or_add_tcPr()
    for m in tcPr.findall(qn("w:tcMar")):
        tcPr.remove(m)
    tcMar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)

    for v in tcPr.findall(qn("w:vAlign")):
        tcPr.remove(v)
    vAlign = OxmlElement("w:vAlign")
    vAlign.set(qn("w:val"), "center")
    tcPr.append(vAlign)

    ps = tc.findall(qn("w:p"))
    for extra in ps[1:]:
        tc.remove(extra)
    p = _first_p(tc)
    Paragraph(p, None).alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p.insert(0, pPr)
    for ind in pPr.findall(qn("w:ind")):
        pPr.remove(ind)
    for sp in pPr.findall(qn("w:spacing")):
        pPr.remove(sp)
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:before"), "0")
    sp.set(qn("w:after"), "0")
    sp.set(qn("w:line"), "240")
    sp.set(qn("w:lineRule"), "auto")
    pPr.append(sp)


def _merge_fullbody_photo(rows, ph):
    """Vertically merge the left photo column (rows 4..29, cell 1) into one tall
    cell and drop the applicant's full-body photo into it."""
    top, bottom, col = 4, 29, 1
    for ri in range(top, bottom + 1):
        cells = rows[ri].findall(qn("w:tc"))
        if col >= len(cells):
            continue
        tc = cells[col]
        _set_vmerge(tc, restart=(ri == top))
        if ri == top:
            _set_cell(tc, ph)
            _format_photo_cell(tc)


def _set_cell(tc, text):
    """Clear the value cell's first paragraph and insert one placeholder run."""
    p = _first_p(tc)
    for r in p.findall(qn("w:r")):
        p.remove(r)
    p.append(_make_run(text))


def _replace_text(tc, text):
    """Replace the cell's existing default text, keeping its run formatting."""
    ts = list(tc.iter(qn("w:t")))
    if not ts:
        _set_cell(tc, text)
        return
    ts[0].text = text
    ts[0].set(qn("xml:space"), "preserve")
    for t in ts[1:]:
        t.text = ""


def _setup_passport_page(doc):
    """Configure page 2 of the template with clean margins and passport placeholder."""
    from docx.shared import Inches, Pt
    if len(doc.sections) > 1:
        sec2 = doc.sections[1]
        sec2.top_margin = Inches(0.5)
        sec2.bottom_margin = Inches(0.5)
        sec2.left_margin = Inches(0.5)
        sec2.right_margin = Inches(0.5)

    if len(doc.paragraphs) > 1:
        p2 = doc.paragraphs[1]
        p2.text = ""
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(4)
        # Jinja conditional block: only renders on page 2 when photo_passport is present
        r_if = _make_run("{% if photo_passport %}", bold=False, size=14)
        p2._element.append(r_if)
        r_title = _make_run("PASSPORT\n", bold=True, size=20, color="1B365D")
        p2._element.append(r_title)
        r_img = _make_run("{{ photo_passport }}", bold=False, size=16)
        p2._element.append(r_img)
        r_endif = _make_run("{% endif %}", bold=False, size=14)
        p2._element.append(r_endif)


def _remove_floating_logo(doc):
    """Remove the floating Future New Logo drawing so it does not overlap with the table logo."""
    for pic in doc.element.xpath(".//pic:pic"):
        for cNvPr in pic.xpath(".//pic:cNvPr"):
            if "Future" in cNvPr.get("descr", "") or "Future" in cNvPr.get("name", ""):
                parent = pic.getparent()
                if parent is not None:
                    parent.remove(pic)
    for shp in doc.element.findall(".//{urn:schemas-microsoft-com:vml}shape"):
        title = shp.attrib.get("{urn:schemas-microsoft-com:office:office}title", "")
        alt = shp.attrib.get("alt", "")
        if "Future" in title or "Future" in alt:
            parent = shp.getparent()
            if parent is not None:
                parent.remove(shp)


def build():
    if not os.path.exists(ORIGINAL):
        raise SystemExit("Missing template/original.docx (the agency form).")
    doc = Document(ORIGINAL)
    tbl = _choice_table(doc)
    if tbl is None:
        raise SystemExit("Could not locate the form table in original.docx")
    rows = tbl.findall(qn("w:tr"))

    # sanity-check that indices still line up with labels
    warns = []
    for (ri, ci), expect in VERIFY.items():
        try:
            got = _cell_text(rows[ri].findall(qn("w:tc"))[ci])
        except IndexError:
            got = "<out of range>"
        if expect.upper() not in got.upper():
            warns.append(f"  R{ri}C{ci}: expected '{expect}', found '{got}'")
    if warns:
        print("WARNING: template layout may have changed:")
        print("\n".join(warns))

    for (ri, ci), (kind, key) in MAP.items():
        cells = rows[ri].findall(qn("w:tc"))
        if ci >= len(cells):
            print(f"  skip R{ri}C{ci} ({key}) - cell missing")
            continue
        tc = cells[ci]
        ph = "{{ %s }}" % key
        if kind == "rep":
            _replace_text(tc, ph)
        else:
            _set_cell(tc, ph)

    # full-body photo: merge the tall left column into one cell
    _merge_fullbody_photo(rows, "{{ photo_full }}")

    # remove floating logo so only the clean cell logo is rendered
    _remove_floating_logo(doc)

    # center and format photo and logo cells
    try:
        _center_cell(rows[1].findall(qn("w:tc"))[1])        # logo box
        _format_photo_cell(rows[1].findall(qn("w:tc"))[2])  # face box
        _format_photo_cell(rows[4].findall(qn("w:tc"))[1])  # full-body box
    except IndexError:
        pass

    # attach passport on page 2
    _setup_passport_page(doc)

    doc.save(OUT)
    print("Template written:", OUT)


if __name__ == "__main__":
    build()
