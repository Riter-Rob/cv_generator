# -*- coding: utf-8 -*-
"""
Excel input sheet: one row per applicant.

- Meta columns tell the tool which passport image / photos to use.
- 'excel' columns are the data NOT on the passport (skills, salary, ...).
- 'override' columns mirror the passport fields; leave them blank to let OCR
  fill them, or type a value to override what OCR reads.

build_input_template(path) -> writes a ready-to-fill .xlsx (with a Guide sheet
and one example row).
read_applicants(path)      -> list of {key: value} dicts (non-empty cells only).
"""
import os
import sys
from datetime import datetime, date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from fields import FIELDS, SKILLS  # noqa: E402

SHEET = "Applicants"

META_COLS = [
    ("output_name",     "Output file name",        "اسم الملف (اختياري)"),
    ("passport_file",   "Passport image file",     "ملف صورة الجواز"),
    ("photo_face_file", "Face photo file",         "ملف صورة الوجه"),
    ("photo_full_file", "Full-body photo file",    "ملف الصورة الكاملة"),
]

EXCEL_COLS = [(k, en, ar) for (k, en, ar, src, d, g) in FIELDS if src == "excel"]
SKILL_COLS = [(k, en + " (YES/NO)", ar) for (k, en, ar, d) in SKILLS]
OVERRIDE_COLS = [(k, en + " *", ar) for (k, en, ar, src, d, g) in FIELDS
                 if src in ("passport", "passport_visual", "derived")]

COLUMNS = META_COLS + EXCEL_COLS + SKILL_COLS + OVERRIDE_COLS
KEYS = [c[0] for c in COLUMNS]

# ---- styling -----------------------------------------------------------------
HDR = Font(bold=True, size=10, color="FFFFFF", name="Arial")
HINT = Font(italic=True, size=8, color="666666", name="Arial")
META_FILL = PatternFill("solid", fgColor="7F7F7F")
EXCEL_FILL = PatternFill("solid", fgColor="C0504D")
SKILL_FILL = PatternFill("solid", fgColor="4F81BD")
OVR_FILL = PatternFill("solid", fgColor="9BBB59")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _fill_for(key):
    keys_meta = {c[0] for c in META_COLS}
    keys_excel = {c[0] for c in EXCEL_COLS}
    keys_skill = {c[0] for c in SKILL_COLS}
    if key in keys_meta:
        return META_FILL
    if key in keys_excel:
        return EXCEL_FILL
    if key in keys_skill:
        return SKILL_FILL
    return OVR_FILL


def build_input_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET

    # header + hint rows
    for ci, (key, en, ar) in enumerate(COLUMNS, start=1):
        c1 = ws.cell(row=1, column=ci, value=key)
        c1.font = HDR
        c1.fill = _fill_for(key)
        c1.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c1.border = BORDER
        c2 = ws.cell(row=2, column=ci, value=f"{en}  |  {ar}")
        c2.font = HINT
        c2.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c2.border = BORDER
        ws.column_dimensions[c1.column_letter].width = max(14, min(24, len(en) + 6))

    # comments on the meta + override headers
    ws.cell(row=1, column=1).comment = Comment(
        "Optional. If blank, the applicant's full name is used for the output file.", "CVGen")
    ws.cell(row=1, column=2).comment = Comment(
        "File name of the passport image inside input/passports/ (e.g. rewda.jpg). "
        "OCR reads it automatically.", "CVGen")
    # find first override column to annotate
    for ci, (key, en, ar) in enumerate(COLUMNS, start=1):
        if key in {c[0] for c in OVERRIDE_COLS}:
            ws.cell(row=1, column=ci).comment = Comment(
                "* Auto-filled from the passport by OCR. Leave blank unless you "
                "want to override the OCR value.", "CVGen")
            break

    ws.freeze_panes = "A3"
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 26

    # example row (generic placeholder - replace with a real applicant).
    # Blank skill/salary/language cells demonstrate the built-in defaults.
    example = {
        "output_name": "SAMPLE APPLICANT",
        "passport_file": "",   # e.g. sample.jpg  (put the image in input/passports/)
        "photo_face_file": "",
        "photo_full_file": "",
        "position": "HOUSE MAID",
        "salary": "",
        "preferred_country": "SAUDI ARABIA",
        "contact_no": "",
        "english_level": "",
        "arabic_level": "",
        "education_level": "HIGH SCHOOL",
        "prev_country": "FIRST TIME",
        "prev_position": "FIRST TIME",
        "prev_period_from": "FIRST TIME",
        "prev_period_to": "FIRST TIME",
        "religion": "MUSLIM",
        "living_town": "",
        "marital_status": "SINGLE",
        "num_children": "0",
        "weight": "",
        "height": "",
        "complexion": "",
        "profile_summary": "",
    }
    for ci, key in enumerate(KEYS, start=1):
        val = example.get(key, "")
        cell = ws.cell(row=3, column=ci, value=val)
        cell.border = BORDER
        cell.font = Font(size=9, name="Arial")

    _add_dropdowns(ws)

    _build_guide(wb)
    wb.save(path)
    return path


def _add_dropdowns(ws):
    """Attach list-validation dropdowns to selected columns (rows 3..500)."""
    lists = {
        "religion": '"MUSLIM,CHRISTIAN"',
    }
    for key, en, ar, d in SKILLS:
        lists[key] = '"YES,NO"'
    for key, formula in lists.items():
        if key not in KEYS:
            continue
        col = get_column_letter(KEYS.index(key) + 1)
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}3:{col}500")


def _build_guide(wb):
    g = wb.create_sheet("Guide")
    headers = ["Column key", "English", "Arabic", "Source", "Notes"]
    for ci, h in enumerate(headers, start=1):
        c = g.cell(row=1, column=ci, value=h)
        c.font = Font(bold=True, color="FFFFFF", name="Arial")
        c.fill = PatternFill("solid", fgColor="404040")
    src_notes = {
        "passport": "Auto from passport MRZ (override optional)",
        "passport_visual": "Auto from passport (printed) - verify/override",
        "derived": "Computed automatically (e.g. age from DOB)",
        "excel": "You must type this",
    }
    r = 2
    for key, en, ar, src, d, grp in FIELDS:
        g.cell(row=r, column=1, value=key)
        g.cell(row=r, column=2, value=en)
        g.cell(row=r, column=3, value=ar)
        g.cell(row=r, column=4, value=src)
        g.cell(row=r, column=5, value=src_notes.get(src, ""))
        r += 1
    for key, en, ar, d in SKILLS:
        g.cell(row=r, column=1, value=key)
        g.cell(row=r, column=2, value=en)
        g.cell(row=r, column=3, value=ar)
        g.cell(row=r, column=4, value="excel")
        g.cell(row=r, column=5, value="Type YES or NO")
        r += 1
    for col, w in zip("ABCDE", (22, 22, 22, 18, 42)):
        g.column_dimensions[col].width = w


# ---- reader ------------------------------------------------------------------
def _norm(val):
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    return str(val).strip()


def read_applicants(path):
    wb = load_workbook(path, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
    header = [_norm(c.value) for c in ws[1]]
    rows = []
    for excel_row_idx, raw in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        # skip the hint row (row 2) and blank rows
        values = [_norm(v) for v in raw]
        if all(v == "" for v in values):
            continue
        # detect hint row: contains ' | ' separators matching our hint format
        if excel_row_idx == 2 and any(" | " in v for v in values):
            continue
        row = {"_excel_row": excel_row_idx}
        for key, val in zip(header, values):
            if not key:
                continue
            if val == "":
                continue
            row[key] = val
        # a row is real only if it has at least a passport_file or a name
        if row.get("passport_file") or row.get("full_name") or row.get("output_name"):
            rows.append(row)
    return rows


if __name__ == "__main__":
    out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                       "input", "applicants.xlsx")
    build_input_template(out)
    print("Excel input template written:", out)
